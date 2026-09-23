import docx, re

doc = docx.Document(r'C:\Users\Sriram\.gemini\antigravity\scratch\online-shopping-dba\DBA-Template.docx')

print("=== CHECKING PARAGRAPHS FOR DRAWINGS ===")
import sys
sys.stdout.reconfigure(encoding='utf-8')
for i, p in enumerate(doc.paragraphs):
    if 'drawing' in p._p.xml:
        embeds = re.findall(r'r:embed="([^"]+)"', p._p.xml)
        print(f"Paragraph {i}: text='{p.text[:40]}' embeds={embeds}")

print("=== CHECKING HEADERS / FOOTERS ===")
for s_idx, s in enumerate(doc.sections):
    for h_name, h in [('header', s.header), ('first_page_header', s.first_page_header), ('footer', s.footer)]:
        if h:
            for p_idx, p in enumerate(h.paragraphs):
                if 'drawing' in p._p.xml:
                    embeds = re.findall(r'r:embed="([^"]+)"', p._p.xml)
                    print(f"Section {s_idx} {h_name} p{p_idx}: text='{p.text[:40]}' embeds={embeds}")

print("=== CHECKING RELATIONSHIPS ===")
for rel_id, rel in doc.part.rels.items():
    if 'image' in rel.target_ref:
        print(f"{rel_id} -> {rel.target_ref}")
