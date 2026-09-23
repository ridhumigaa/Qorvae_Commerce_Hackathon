import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

template_path = r'C:\Users\Sriram\.gemini\antigravity\scratch\online-shopping-dba\DBA-Template.docx'
output_path = r'C:\Users\Sriram\.gemini\antigravity\scratch\online-shopping-dba\Online_Shopping_Database_Report.docx'

doc = docx.Document(template_path)

# Verify that headers and logos exist
print("Sections count:", len(doc.sections))
for s_idx, s in enumerate(doc.sections):
    rels = [rel.target_ref for rel in s.header.part.rels.values() if 'image' in rel.target_ref]
    print(f"Section {s_idx} header images:", rels)

# 1. Update Title (Paragraph 0)
doc.paragraphs[0].text = "Online Shopping Database (E-Commerce Management System)"
if doc.paragraphs[0].runs:
    doc.paragraphs[0].runs[0].font.bold = True
    doc.paragraphs[0].runs[0].font.size = Pt(20)
    doc.paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 23, 42)

# 2. Update Table 0: Student Details
t0 = doc.tables[0]
student_details = [
    ("Student Name", "Ajay Raja"),
    ("Register Number", "732119104701"),
    ("Department", "CSE"),
    ("ID", "8da3a429c5260608bcc40cf1e7433ef3"),
    ("College", "Nandha engineering college, Erode"),
    ("Date of Submission", "23/09/2026"),
    ("GitHub Repo", "https://github.com/AjayGitPersonal/port-final.git"),
    ("Deployment Link", "http://localhost:5000 (Full-Stack Storefront & Oracle DBA Console)")
]

for i, (label, val) in enumerate(student_details):
    if i < len(t0.rows):
        t0.rows[i].cells[0].text = label
        t0.rows[i].cells[1].text = val
        t0.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
        t0.rows[i].cells[0].paragraphs[0].runs[0].font.size = Pt(10)
        t0.rows[i].cells[1].paragraphs[0].runs[0].font.size = Pt(10)

# 3. Update Table 1: Tables Used
t1 = doc.tables[1]
while len(t1.rows) > 1:
    t1._tbl.remove(t1.rows[1]._tr)

tables_info = [
    ('CATEGORIES', 'Stores product classification categories, supporting hierarchical category trees'),
    ('CUSTOMERS', 'Stores customer profiles, contact numbers, email addresses, and delivery coordinates'),
    ('PRODUCTS', 'Stores physical product catalog, SKUs, sales prices, and live warehouse stock levels'),
    ('ORDERS', 'Tracks financial checkout transactions, timestamps, payment methods, and fulfillment lifecycle'),
    ('ORDER_ITEMS', 'Associative table storing individual line items, quantities, and snapshot prices per order'),
    ('AUDIT_LOGS', 'Change Data Capture (CDC) table recording DBA audit trails on inventory mutations')
]
for name, desc in tables_info:
    row = t1.add_row()
    row.cells[0].text = name
    row.cells[1].text = desc
    row.cells[0].paragraphs[0].runs[0].font.bold = True
    row.cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)
    row.cells[1].paragraphs[0].runs[0].font.size = Pt(9.5)

# 4. Update Table 2: Technologies Used
t2 = doc.tables[2]
while len(t2.rows) > 1:
    t2._tbl.remove(t2.rows[1]._tr)

tech_info = [
    ('Oracle Database 21c XE', 'Enterprise relational database engine hosting pluggable database XEPDB1'),
    ('Oracle SQL*Plus', 'Command-line database administration, tablespace management, and script execution'),
    ('SQL & PL/SQL', 'DDL/DML definitions, views, triggers, and atomic inventory checkout packages'),
    ('Node.js & Express.js', 'High-performance REST API backend handling connection pooling and transactions'),
    ('node-oracledb (Thin Mode)', 'Official native Oracle driver communicating via TCP port 1521 without heavy client binaries'),
    ('React 18 & Vite', 'Modern single-page storefront, customer order history UI, and DBA command center'),
    ('Tailwind CSS', 'Responsive styling using the enterprise Qorvae golden-amber palette')
]
for tech, purpose in tech_info:
    row = t2.add_row()
    row.cells[0].text = tech
    row.cells[1].text = purpose
    row.cells[0].paragraphs[0].runs[0].font.bold = True
    row.cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)
    row.cells[1].paragraphs[0].runs[0].font.size = Pt(9.5)

# 5. Update Table 3: Constraints and Relationships
t3 = doc.tables[3]
while len(t3.rows) > 1:
    t3._tbl.remove(t3.rows[1]._tr)

constraints_info = [
    ('CATEGORIES', 'category_id', 'parent_category_id -> CATEGORIES(category_id)', 'category_name NOT NULL UNIQUE, is_active CHECK (0, 1)'),
    ('CUSTOMERS', 'customer_id', 'None', 'first_name NOT NULL, last_name NOT NULL, email NOT NULL UNIQUE, status CHECK (ACTIVE, SUSPENDED, INACTIVE)'),
    ('PRODUCTS', 'product_id', 'category_id -> CATEGORIES(category_id)', 'sku NOT NULL UNIQUE, name NOT NULL, price CHECK (>= 0), stock_quantity CHECK (>= 0)'),
    ('ORDERS', 'order_id', 'customer_id -> CUSTOMERS(customer_id)', 'order_date NOT NULL, total_amount CHECK (>= 0), status CHECK (PENDING, PROCESSING, SHIPPED, DELIVERED, CANCELLED)'),
    ('ORDER_ITEMS', 'order_item_id', 'order_id -> ORDERS(order_id), product_id -> PRODUCTS(product_id)', 'quantity CHECK (> 0), unit_price CHECK (>= 0), subtotal CHECK (>= 0), ON DELETE CASCADE on order_id')
]
for tbl, pk, fk, other in constraints_info:
    row = t3.add_row()
    row.cells[0].text = tbl
    row.cells[1].text = pk
    row.cells[2].text = fk
    row.cells[3].text = other
    row.cells[0].paragraphs[0].runs[0].font.bold = True
    for c in row.cells:
        c.paragraphs[0].runs[0].font.size = Pt(8.5)

# 6. Deep Content Replacements across all Paragraphs
replacements = [
    ('Vehicle Service Management Database', 'Online Shopping Database (E-Commerce Management System)'),
    ('vehicle service center', 'online shopping e-commerce platform'),
    ('vehicle service management', 'online shopping management'),
    ('vehicle service records', 'customer order history and purchasing receipts'),
    ('vehicle service', 'online shopping'),
    ('vehicle_service_db', 'ECOMMERCE_DBA (Oracle 21c PDB: XEPDB1)'),
    ('MySQL Workbench', 'Oracle SQL*Plus & Enterprise DBA Console'),
    ('MySQL', 'Oracle Database 21c'),
    ('Customer, Vehicle, and Service tables', 'Categories, Customers, Products, Orders, and Order Items tables'),
    ('Customer, Vehicle, and Service', 'Categories, Customers, Products, Orders, and Order Items'),
    ('Customer -> Vehicle and Vehicle -> Service', 'Customers -> Orders -> Order Items -> Products -> Categories'),
    ('Customer -> Vehicle', 'Customers -> Orders'),
    ('Vehicle -> Service', 'Orders -> Order Items'),
    ('Customer → Vehicle and Vehicle → Service', 'Customers → Orders → Order Items → Products → Categories'),
    ('Customer → Vehicle', 'Customers → Orders'),
    ('Vehicle → Service', 'Orders → Order Items'),
    ('vehicle information', 'product catalog and warehouse inventory information'),
    ('vehicle details', 'product specifications, SKUs, and stock quantities'),
    ('service records', 'order transactions and itemized line items'),
    ('services performed on those vehicles', 'products ordered by customers'),
    ('service history', 'customer order history'),
    ('service costs', 'order totals and line-item subtotals'),
    ('service cost', 'order total amount'),
    ('service-related costs', 'order billing costs'),
    ('vehicle_id', 'order_id'),
    ('vehicle_number', 'product_sku'),
    ('vehicle_type', 'category_name'),
    ('vehicle_model', 'product_name'),
    ('service_id', 'order_item_id'),
    ('service_type', 'payment_method'),
    ('service_status', 'order_status'),
    ('service_date', 'order_date'),
    ('Oil Change', 'Laptop Purchase (PROD-LAP-001)'),
    ('Brake Service', 'ANC Headphones (PROD-AUD-001)'),
    ('General Service', 'Smart Controller (PROD-IOT-001)'),
    ('AC Service', 'UltraVision 4K Monitor'),
    ('Engine Service', 'DBA Mechanical Keyboard'),
    ('Tyre Replacement', 'Wireless Ergonomic Mouse'),
    ('TN38AB1234', 'PROD-LAP-001'),
    ('TN37CD5678', 'PROD-AUD-001'),
    ('TN33EF9012', 'PROD-IOT-001'),
    ('TN39GH3456', 'PROD-MON-001'),
    ('TN42JK7890', 'PROD-KBD-001'),
    ('TN40LM2468', 'PROD-MOU-001'),
    ('TN45NP1357', 'PROD-ACC-001'),
    ('Hyundai i20', 'Apex Pro 16" Creator Laptop'),
    ('Royal Enfield Classic 350', 'AcousticMax Wireless ANC Headphones'),
    ('Maruti Suzuki Swift', 'SmartHub Quantum Controller'),
    ('Tata Nexon', 'UltraVision 4K Gaming Monitor'),
    ('Honda City', 'Keychron K2 Mechanical Keyboard'),
    ('Yamaha R15', 'Logitech MX Master 3S'),
    ('Toyota Glanza', 'Thunderbolt 4 Pro Dock')
]

for p in doc.paragraphs:
    orig_text = p.text
    new_text = orig_text
    for old_val, new_val in replacements:
        if old_val in new_text:
            new_text = new_text.replace(old_val, new_val)
    if new_text != orig_text:
        p.text = new_text

# Save modified Word document over both paths
doc.save(output_path)
doc.save(template_path)
print("Successfully updated both Online_Shopping_Database_Report.docx and DBA-Template.docx with Ajay Raja's details and preserved header logos!")
